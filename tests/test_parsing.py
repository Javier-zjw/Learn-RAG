"""
文档解析的测试（结构切分的测试在 test_chunking.py）。

样例文件都在临时目录里现场生成，不提交二进制文件；
MinerU、Docling、VLM OCR 这类重型依赖一律用假对象或 mock，不依赖 GPU 和网络。
"""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from learn_rag.core.registry import registry
from learn_rag.core.types import Document, Element
from learn_rag.ingest.structure import StructureChunker
from learn_rag.parsing.external import MinerUParser, docling_to_elements, mineru_middle_to_elements
from learn_rag.parsing.markdown import markdown_to_elements, table_to_markdown
from learn_rag.parsing.source import FileSource
from learn_rag.parsing.web import HtmlParser


def _kinds(elements: list[Element]) -> list[str]:
    return [e.kind for e in elements]


class TestMarkdown(unittest.TestCase):
    def test_blocks(self):
        text = (
            "# 总则\n\n第一段，\n同一段的第二行。\n\n## 计费\n\n"
            "| 项目 | 价格 |\n| --- | --- |\n| A | 1 |\n\n"
            "```\nprint(1)\n```\n\n$$\nE=mc^2\n$$\n\n![架构图](a.png)\n\n"
            "<table><tr><td>合并</td></tr>\n</table>\n"
        )
        elements = markdown_to_elements(text, page=3)
        self.assertEqual(_kinds(elements), ["heading", "text", "heading", "table", "code", "formula", "image", "table"])
        self.assertEqual((elements[2].text, elements[2].level), ("计费", 2))
        self.assertEqual(elements[1].text, "第一段，\n同一段的第二行。")
        self.assertEqual(elements[5].text, "E=mc^2")
        self.assertEqual(elements[6].extra["src"], "a.png")
        self.assertTrue(all(e.page == 3 for e in elements))

    def test_table_to_markdown(self):
        md = table_to_markdown([["名称", "说明"], ["a|b", "多\n行"], ["c"], ["", None]])
        self.assertEqual(md.splitlines(), ["| 名称 | 说明 |", "| --- | --- |", "| a\\|b | 多 行 |", "| c |  |"])

    def test_normalize_table(self):
        from learn_rag.parsing.markdown import normalize_table

        # 没有合并单元格：带缩进和 <p> 的 HTML 转成 Markdown，日期去掉零点时间，实体和 <br> 正确处理
        pretty = ("<table>\n  <tr>\n    <th><p>订单</p></th>\n    <th><p>日期</p></th>\n  </tr>\n"
                  "  <tr>\n    <td><p>A|1</p></td>\n    <td><p>2025-10-08 00:00:00</p></td>\n  </tr>\n"
                  "  <tr>\n    <td>含&quot;引号&quot;</td>\n    <td>第一行<br>第二行</td>\n  </tr>\n</table>")
        self.assertEqual(normalize_table(pretty).splitlines(),
                         ["| 订单 | 日期 |", "| --- | --- |", "| A\\|1 | 2025-10-08 |", '| 含"引号" | 第一行 第二行 |'])
        # 有合并单元格：保留 HTML，但只留结构
        merged = '<table border="1">\n <tr><th rowspan="2"><strong>版本</strong></th><th colspan="2">价格</th></tr>\n' \
                 ' <tr><td>月费</td><td>年费 &amp; 折扣</td></tr></table>'
        self.assertEqual(normalize_table(merged),
                         '<table><tr><td rowspan="2">版本</td><td colspan="2">价格</td></tr>'
                         '<tr><td>月费</td><td>年费 &amp; 折扣</td></tr></table>')
        self.assertEqual(normalize_table("| a |\n| --- |"), "| a |\n| --- |")

    def test_document_from_elements_renders_markdown(self):
        doc = Document.from_elements("d", [Element("heading", "标题", level=2), Element("text", "正文")])
        self.assertEqual(doc.text, "## 标题\n\n正文")
        self.assertEqual(len(doc.elements), 2)


class TestHtml(unittest.TestCase):
    def test_structure_and_noise(self):
        html = """<html><head><style>p{}</style><script>var x=1;</script></head><body>
        <nav>首页 | 关于</nav>
        <h1>产品手册</h1>
        <p>第一行<br>第二行
           仍是第二行</p>
        <h2>价格</h2>
        <table><tr><th>型号</th><th>价格</th></tr><tr><td>A1</td><td>100</td></tr></table>
        <pre>line1
line2</pre>
        <footer>版权所有</footer></body></html>"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.html"
            path.write_text(html, encoding="utf-8")
            elements = HtmlParser().parse(path)
        self.assertEqual(_kinds(elements), ["heading", "text", "heading", "table", "code"])
        self.assertEqual(elements[1].text, "第一行\n第二行 仍是第二行")
        self.assertIn("| A1 | 100 |", elements[3].text)
        self.assertEqual(elements[4].text, "line1\nline2")
        self.assertNotIn("版权所有", " ".join(e.text for e in elements))


class TestOffice(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_keeps_table_position(self):
        try:
            import docx
        except ImportError:
            self.skipTest("未安装 python-docx")
        from learn_rag.parsing.office import DocxParser

        d = docx.Document()
        d.add_heading("报销制度", level=1)
        d.add_paragraph("适用于全体员工。")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "城市", "上限"
        table.cell(1, 0).text, table.cell(1, 1).text = "一线", "600"
        d.add_heading("附则", level=2)
        path = self.dir / "a.docx"
        d.save(path)

        elements = DocxParser().parse(path)
        self.assertEqual(_kinds(elements), ["heading", "text", "table", "heading"])
        self.assertEqual((elements[0].level, elements[3].level), (1, 2))
        self.assertIn("| 一线 | 600 |", elements[2].text)

    def test_pptx_slides_tables_notes(self):
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            self.skipTest("未安装 python-pptx")
        from learn_rag.parsing.office import PptxParser

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])   # 只有标题的版式
        slide.shapes.title.text = "季度营收"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "区域", "营收"
        table.cell(1, 0).text, table.cell(1, 1).text = "华东", "480"
        slide.notes_slide.notes_text_frame.text = "华东贡献最大"
        prs.slides.add_slide(prs.slide_layouts[6])            # 空白页：用页码作为标题
        path = self.dir / "a.pptx"
        prs.save(path)

        elements = PptxParser().parse(path)
        self.assertEqual(_kinds(elements), ["heading", "table", "text", "heading"])
        self.assertEqual(elements[0].text, "季度营收")
        self.assertIn("| 华东 | 480 |", elements[1].text)
        self.assertEqual(elements[2].text, "备注：华东贡献最大")
        self.assertEqual((elements[3].text, elements[3].page), ("第 2 页", 2))

    def test_xlsx_drops_empty_rows_and_columns(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("未安装 openpyxl")
        from learn_rag.parsing.office import XlsxParser

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "订单"
        ws["A1"], ws["C1"] = "区域", "数量"
        ws["A2"], ws["C2"] = "华东", 120
        ws["A5"] = "合计"
        ws["C5"] = "=C2"
        path = self.dir / "a.xlsx"
        wb.save(path)

        elements = XlsxParser().parse(path)
        self.assertEqual(_kinds(elements), ["heading", "table"])
        self.assertEqual(elements[0].text, "订单")
        self.assertEqual(elements[1].text.splitlines()[0], "| 区域 | 数量 |")
        self.assertEqual(len(elements[1].text.splitlines()), 4)   # 表头 + 分隔线 + 两行数据


class TestPdf(unittest.TestCase):
    def setUp(self):
        try:
            import pymupdf  # noqa: F401
        except ImportError:
            self.skipTest("未安装 pymupdf")
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "a.pdf"

    def tearDown(self):
        self.tmp.cleanup()

    def _make_pdf(self, scanned_last_page: bool = False):
        import pymupdf

        pdf = pymupdf.open()
        for number in (1, 2, 3):
            page = pdf.new_page()
            page.insert_text((72, 30), "Company Confidential", fontsize=9)       # 页眉
            page.insert_text((290, 820), f"- {number} -", fontsize=9)           # 页脚页码
            if number == 1:
                page.insert_text((72, 100), "Travel Policy", fontsize=20)
                page.insert_text((72, 140), "Hotels in tier one cities are capped at 600 per night and", fontsize=11)
            elif number == 2:
                page.insert_text((72, 100), "other cities at 400 per night.", fontsize=11)
                page.insert_text((72, 140), "Meals", fontsize=15)
                page.insert_text((72, 170), "Meal allowance is 120 per day for every employee.", fontsize=11)
            elif not scanned_last_page:
                page.insert_text((72, 100), "Receipts must be submitted within thirty days.", fontsize=11)
        pdf.save(self.path)

    def test_headings_noise_and_cross_page_merge(self):
        from learn_rag.parsing.pdf import PdfParser

        self._make_pdf()
        elements = PdfParser().parse(self.path)
        texts = [e.text for e in elements]
        self.assertEqual(_kinds(elements), ["heading", "text", "heading", "text", "text"])
        self.assertEqual((elements[0].text, elements[0].level), ("Travel Policy", 1))
        self.assertEqual((elements[2].text, elements[2].level), ("Meals", 2))
        self.assertIn("600 per night and other cities at 400 per night.", texts[1])
        self.assertEqual((elements[1].page, elements[1].extra.get("page_end")), (1, 2))
        self.assertFalse(any("Confidential" in t or t.startswith("-") for t in texts))
        self.assertFalse(any("size" in e.extra or "margin" in e.extra for e in elements))

    def test_scanned_page_goes_to_ocr(self):
        from learn_rag.parsing.pdf import PdfParser

        class FakeOcr:
            def __init__(self):
                self.pages = []

            def recognize(self, image: bytes, *, page=None, mime="image/png"):
                self.pages.append(page)
                self.assertion = image.startswith(b"\x89PNG")
                return [Element("text", "OCR 识别出的内容。", page=page)]

        self._make_pdf(scanned_last_page=True)
        ocr = FakeOcr()
        elements = PdfParser(ocr=ocr, min_chars_per_page=40).parse(self.path)
        self.assertEqual(ocr.pages, [3])
        self.assertTrue(ocr.assertion)
        self.assertEqual((elements[-1].text, elements[-1].page), ("OCR 识别出的内容。", 3))

    def test_table_detected(self):
        import pymupdf

        from learn_rag.parsing.pdf import PdfParser

        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 80), "Price list for all products in the current catalogue.", fontsize=11)
        x = [72, 222, 372]
        y = [120, 145, 170]
        for yy in y:
            page.draw_line((x[0], yy), (x[-1], yy))
        for xx in x:
            page.draw_line((xx, y[0]), (xx, y[-1]))
        for row, (a, b) in enumerate([("Model", "Price"), ("A1", "100")]):
            page.insert_text((x[0] + 5, y[row] + 17), a, fontsize=11)
            page.insert_text((x[1] + 5, y[row] + 17), b, fontsize=11)
        pdf.save(self.path)

        elements = PdfParser().parse(self.path)
        self.assertEqual(_kinds(elements), ["text", "table"])
        self.assertIn("| A1 | 100 |", elements[1].text)


class TestVlmOcr(unittest.TestCase):
    def test_sends_image_and_parses_markdown(self):
        from learn_rag.parsing.ocr import VlmOcrParser

        sent = {}

        def fake_chat(messages, **_):
            sent["messages"] = messages
            return "```markdown\n# 发票\n\n金额：100 元\n```"

        parser = VlmOcrParser(model="ocr-model", base_url="http://localhost/v1", api_key="k")
        with patch.object(parser.llm, "chat", side_effect=fake_chat):
            elements = parser.recognize(b"\x89PNGfake", page=2)
        content = sent["messages"][0]["content"]
        self.assertTrue(content[0]["image_url"]["url"].startswith("data:image/png;base64,"))
        self.assertEqual(_kinds(elements), ["heading", "text"])
        self.assertEqual((elements[1].text, elements[1].page, elements[1].extra["ocr"]), ("金额：100 元", 2, True))

    def test_requires_model(self):
        from learn_rag.parsing.ocr import VlmOcrParser

        with patch.dict("os.environ", {"OCR_MODEL": ""}):
            with self.assertRaises(ValueError):
                VlmOcrParser()


class TestExternalAdapters(unittest.TestCase):
    MIDDLE = {
        "schema": "docvortex.middle",
        "pages": [
            {
                "page_idx": 0,
                "blocks": [
                    {"type": "aside_text", "content": [{"type": "text", "content": "内部资料"}]},
                    {"type": "doc_title", "level": 1, "bbox": [0.1, 0.1, 0.8, 0.2],
                     "content": [{"type": "text", "content": "企业手册"}]},
                    {"type": "paragraph_title", "level": 2,
                     "content": [{"type": "text", "content": "第一章 总则"}]},
                    {"type": "text", "bbox": [0.1, 0.30, 0.9, 0.32],
                     "content": [{"type": "text", "content": "正文开头，"}]},
                    {"type": "table", "content": [
                        {"type": "table_caption", "content": [{"type": "text", "content": "表1 价格"}]},
                        {"type": "table_body", "image_path": "images/table.jpg",
                         "content": "<table><tr><td>A</td></tr></table>"},
                    ]},
                    {"type": "image", "content": [
                        {"type": "image_body", "image_path": "images/pic.jpg", "content": ""},
                        {"type": "image_caption", "content": [{"type": "text", "content": "图1 架构"}]},
                    ]},
                    {"type": "chart", "content": [
                        {"type": "chart_body", "image_path": "images/chart.jpg", "content": ""},
                    ]},
                    {"type": "equation", "content": "E=mc^2"},
                    {"type": "code", "guess_lang": "python", "content": [
                        {"type": "code_body", "content": "print(1)"},
                    ]},
                    {"type": "list", "content": [
                        {"type": "text", "content": [{"type": "text", "content": "要点一"}]},
                        {"type": "text", "content": [{"type": "text", "content": "要点二"}]},
                    ]},
                    {"type": "page_number", "content": [{"type": "text", "content": "1"}]},
                ],
            },
            {
                "page_idx": 1,
                "blocks": [
                    {"type": "text", "bbox": [0.1, 0.05, 0.6, 0.07],
                     "content": [{"type": "text", "content": "接着延续。"}]},
                ],
            },
        ],
    }
    IMAGES = {
        "images/table.jpg": b"table-bytes",
        "images/pic.jpg": b"image-bytes",
        "images/chart.jpg": b"chart-bytes",
    }

    @staticmethod
    def _write_mineru_zip(path: Path, middle=None, images=None):
        with ZipFile(path, "w") as archive:
            archive.writestr("middle_json.json", json.dumps(middle or TestExternalAdapters.MIDDLE))
            for name, data in (images if images is not None else TestExternalAdapters.IMAGES).items():
                archive.writestr(name, data)

    def test_mineru_middle_blocks(self):
        elements = mineru_middle_to_elements(self.MIDDLE)
        self.assertEqual(_kinds(elements),
                         ["heading", "heading", "text", "table", "image", "image", "formula", "code", "text", "text"])
        self.assertEqual((elements[0].level, elements[0].page, elements[0].bbox), (1, 1, [0.1, 0.1, 0.8, 0.2]))
        self.assertEqual(elements[1].level, 2)
        self.assertEqual((elements[2].text, elements[2].page), ("正文开头，", 1))
        self.assertEqual(elements[3].text, "表1 价格\n| A |\n| --- |")          # 没有合并单元格的 HTML 表格转成 Markdown
        self.assertEqual((elements[4].text, elements[4].extra["mineru_type"]), ("图1 架构", "image"))
        self.assertEqual((elements[5].text, elements[5].extra["mineru_type"]), ("[图表]", "chart"))
        self.assertEqual(elements[6].text, "E=mc^2")
        self.assertEqual(elements[7].text, "print(1)")
        self.assertEqual(elements[8].text, "要点一\n要点二")
        # 第 1 页最后是列表，不是"正文开头，"，所以第 2 页的文字不能接到它后面
        self.assertEqual((elements[9].text, elements[9].page), ("接着延续。", 2))

    def test_mineru_runs_cli_and_reads_zip(self):
        def fake_run(cmd, **_):
            self._write_mineru_zip(Path(cmd[cmd.index("-o") + 1]))
            return SimpleNamespace(returncode=0)

        with patch("learn_rag.parsing.external.subprocess.run", side_effect=fake_run) as run:
            elements = MinerUParser().parse(Path("a.pdf"))
        cmd = run.call_args[0][0]
        # 装了 MinerU 的机器上命令是完整路径（PATH 或 MINERU_PYTHON_ENV/bin 下找到的）
        self.assertEqual((Path(cmd[0]).name, *cmd[1:3]), ("mineru-kit", "parse", "a.pdf"))
        self.assertIn("zip", cmd)
        self.assertIn("standard", cmd)
        self.assertEqual(len(elements), 10)

    def test_mineru_reads_env_file(self):
        """.mineru.env 与启动脚本共用：文件里的值优先于 shell 环境变量和解析器默认值。"""
        import os

        with tempfile.TemporaryDirectory() as tmp:
            env_file = Path(tmp) / ".mineru.env"
            env_file.write_text(
                "# 注释行\n"
                f"export MINERU_HOME={tmp}/models\n"
                "MINERU_MODEL_VLM_SERVER_URL='http://10.0.0.8:30000'  \n"
                "MINERU_MODEL_VLM_ENGINE=vllm  # 行尾注释\n",
                encoding="utf-8",
            )
            with patch.dict(os.environ, {"MINERU_MODEL_VLM_ENGINE": "llama-cpp", "MINERU_MODEL_SOURCE": "local"}):
                parser = MinerUParser(env_file=str(env_file))
            self.assertEqual(parser.models_dir, Path(tmp) / "models")
            self.assertTrue(parser.models_dir.is_dir())
            self.assertEqual(parser.env["MINERU_MODEL_VLM_SERVER_URL"], "http://10.0.0.8:30000")
            self.assertEqual(parser.env["MINERU_MODEL_VLM_ENGINE"], "vllm")     # 文件优先于 shell
            self.assertEqual(parser.env["MINERU_MODEL_SOURCE"], "local")        # shell 优先于默认值
            self.assertEqual(parser.env["MINERU_MODEL_SMALL_BACKEND"], "onnx")  # 都没有时用默认值

    def test_mineru_keeps_raw_zip(self):
        def fake_run(cmd, **_):
            self._write_mineru_zip(Path(cmd[cmd.index("-o") + 1]), images={})
            return SimpleNamespace(returncode=0)

        with tempfile.TemporaryDirectory() as tmp, \
             patch("learn_rag.parsing.external.subprocess.run", side_effect=fake_run):
            MinerUParser(raw_dir=tmp).parse(Path("a.pdf"))
            with ZipFile(Path(tmp) / "a.pdf.zip") as archive:
                self.assertIn("middle_json.json", archive.namelist())

    def test_mineru_replays_saved_result(self):
        """replay_dir 里有保存的结果时直接读取，不调用 MinerU；没有时照常调用。"""
        with tempfile.TemporaryDirectory() as tmp:
            self._write_mineru_zip(Path(tmp) / "a.pdf.zip", images={})
            with patch("learn_rag.parsing.external.subprocess.run", side_effect=AssertionError("不应调用 MinerU")):
                elements = MinerUParser(replay_dir=tmp).parse(Path("a.pdf"))
            self.assertEqual(elements[0].text, "企业手册")
            with patch("learn_rag.parsing.external.subprocess.run", side_effect=FileNotFoundError):
                with self.assertRaisesRegex(RuntimeError, "mineru"):
                    MinerUParser(replay_dir=tmp).parse(Path("other.pdf"))

    def test_mineru_office_files_skip_tier(self):
        """Office 等格式在 MinerU 里固定走 flash，显式传 --tier 会报错，所以不能带档位参数。"""
        def fake_run(cmd, **_):
            self._write_mineru_zip(Path(cmd[cmd.index("-o") + 1]), images={})
            return SimpleNamespace(returncode=0)

        for name in ("a.docx", "a.pptx", "a.xlsx"):
            with patch("learn_rag.parsing.external.subprocess.run", side_effect=fake_run) as run:
                MinerUParser(image_analysis=False).parse(Path(name))
            cmd = run.call_args[0][0]
            self.assertNotIn("--tier", cmd)
            self.assertNotIn("--ocr-mode", cmd)
            self.assertNotIn("--disable-image-analysis", cmd)

    def test_mineru_missing_command(self):
        with patch("learn_rag.parsing.external.subprocess.run", side_effect=FileNotFoundError):
            with self.assertRaisesRegex(RuntimeError, "mineru"):
                MinerUParser().parse(Path("a.pdf"))

    def test_mineru_missing_command(self):
        with patch("learn_rag.parsing.external.subprocess.run", side_effect=FileNotFoundError):
            with self.assertRaisesRegex(RuntimeError, "mineru"):
                MinerUParser().parse(Path("a.pdf"))

    def test_mineru_loads_weights_from_local_dir(self):
        """模型权重必须从项目内目录加载：MINERU_HOME 显式指向 models_dir。"""
        import os

        seen = {}

        def fake_run(cmd, **kwargs):
            seen["env"] = dict(kwargs["env"])
            self._write_mineru_zip(Path(cmd[cmd.index("-o") + 1]), images={})
            return SimpleNamespace(returncode=0)

        with tempfile.TemporaryDirectory() as tmp:
            # 清掉外部环境里的同名变量，验证解析器自己注入了正确的值
            with patch.dict(os.environ, {}, clear=False):
                for key in ("MINERU_HOME", "MINERU_MODEL_SOURCE", "MINERU_MODEL_SMALL_BACKEND",
                            "MINERU_MODEL_VLM_ENGINE", "MINERU_MODEL_VLM_SERVER_URL"):
                    os.environ.pop(key, None)
                with patch("learn_rag.parsing.external.subprocess.run", side_effect=fake_run):
                    MinerUParser(models_dir=tmp, env_file=None).parse(Path("a.pdf"))
            self.assertEqual(seen["env"]["MINERU_HOME"], str(Path(tmp).resolve()))
            self.assertEqual(seen["env"]["MINERU_MODEL_SOURCE"], "modelscope")
            self.assertEqual(seen["env"]["MINERU_MODEL_SMALL_BACKEND"], "onnx")
            self.assertEqual(seen["env"]["MINERU_MODEL_VLM_ENGINE"], "llama-cpp")
            self.assertEqual(seen["env"]["MINERU_MODEL_VLM_SERVER_URL"], "http://127.0.0.1:30000")

    def test_mineru_default_weights_dir_in_project(self):
        root = Path(__file__).resolve().parents[1]
        with patch.dict("os.environ", {}, clear=False):
            import os

            os.environ.pop("MINERU_HOME", None)
            parser = MinerUParser(env_file=None)
        self.assertEqual(parser.models_dir, root / "mineru_model_weight")
        self.assertTrue(parser.models_dir.is_dir())

    def test_docling_items(self):
        def item(label, text="", page=1, **extra):
            box = SimpleNamespace(l=0, t=1, r=2, b=3)
            return SimpleNamespace(label=SimpleNamespace(value=label), text=text,
                                   prov=[SimpleNamespace(page_no=page, bbox=box)], **extra)

        table = item("table", export_to_markdown=lambda doc: "| a |\n| --- |", caption_text=lambda doc: "表1")
        items = [
            item("page_header", "页眉"), item("title", "手册"), item("section_header", "安装", level=1),
            item("text", "步骤一。"), table, item("caption", "表1"),
            item("picture", caption_text=lambda doc: "流程图"), item("list_item", "要点", page=2),
        ]
        doc = SimpleNamespace(iterate_items=lambda: [(i, 0) for i in items])
        elements = docling_to_elements(doc)
        self.assertEqual(_kinds(elements), ["heading", "heading", "text", "table", "image", "text"])
        self.assertEqual([e.level for e in elements[:2]], [1, 2])
        self.assertEqual(elements[3].text, "表1\n| a |\n| --- |")
        self.assertEqual((elements[5].page, elements[5].bbox), (2, [0, 1, 2, 3]))


def _t(text: str, bbox: list[float] | None = None, kind: str = "text", **extra) -> dict:
    """构造一个 MinerU 文字块（内容是 InlineSpan 列表，和真实输出一致）。"""
    block = {"type": kind, "content": [{"type": "text", "content": text}], **extra}
    if bbox is not None:
        block["bbox"] = bbox
    return block


def _middle(*pages: list[dict]) -> dict:
    return {"schema": "docvortex.middle", "pages": [{"page_idx": i, "blocks": list(b)} for i, b in enumerate(pages)]}


class TestMinerUConversion(unittest.TestCase):
    """按 samples/parsing_results/mineru_raw 里真实 MinerU 4.x 输出的形态构造的用例。"""

    def test_footnotes_and_speaker_notes_are_kept(self):
        """PDF 的页脚注释和 PPT 的演讲者备注在 MinerU 里都是 page_footnote，属于正文内容。"""
        elements = mineru_middle_to_elements(_middle(
            [_t("正文。", [0.1, 0.1, 0.9, 0.12]), _t("¹ 客户留存率 = 年末客户数 / 年初客户数。", [0.1, 0.9, 0.6, 0.92], "page_footnote"),
             _t("第 1 页", [0.45, 0.95, 0.55, 0.96], "page_number")],
            [_t("讲解要点：先说结论。", kind="page_footnote")],
        ))
        self.assertEqual([e.text for e in elements], ["正文。", "¹ 客户留存率 = 年末客户数 / 年初客户数。", "讲解要点：先说结论。"])
        self.assertTrue(all(e.extra.get("footnote") for e in elements[1:]))

    def test_native_chart_becomes_data_table(self):
        """PPT / Excel 的原生图表：chart_body 是图表数据转成的 HTML 表格，要作为表格保留下来。"""
        chart = {"type": "chart", "content": [
            {"type": "chart_caption", "content": [{"type": "text", "content": "各区域季度营收"}]},
            {"type": "chart_body", "content": "<table>\n  <tr>\n    <th><p>区域</p></th>\n    <th><p>Q1</p></th>\n  </tr>\n"
                                              "  <tr>\n    <td><p>华东</p></td>\n    <td><p>320</p></td>\n  </tr>\n</table>"},
        ]}
        empty_chart = {"type": "chart", "content": [{"type": "chart_body", "content": "", "image_path": "images/c.jpg"}]}
        elements = mineru_middle_to_elements(_middle([chart, empty_chart]))
        self.assertEqual(_kinds(elements), ["table", "image"])
        self.assertEqual(elements[0].text, "各区域季度营收\n| 区域 | Q1 |\n| --- | --- |\n| 华东 | 320 |")
        self.assertEqual((elements[1].text, elements[1].extra["image_path"]), ("[图表]", "images/c.jpg"))

    def test_image_description_but_not_file_name(self):
        def image(content: str) -> dict:
            return {"type": "image", "content": [{"type": "image_body", "content": content, "image_path": "images/i.png"}]}

        elements = mineru_middle_to_elements(_middle([image("image.png"), image("一张展示四个审批环节的流程图")]))
        self.assertEqual([e.text for e in elements], ["[图片]", "一张展示四个审批环节的流程图"])

    def test_visual_footnote_becomes_paragraph_and_joins_next_line(self):
        """MinerU 会把图后面的正文段落识别成图注；拆成独立段落后，还要和下一行接起来。"""
        chart = {"type": "chart", "bbox": [0.25, 0.5, 0.75, 0.7], "content": [
            {"type": "chart_body", "content": "", "bbox": [0.25, 0.5, 0.75, 0.7]},
            {"type": "chart_caption", "bbox": [0.4, 0.72, 0.6, 0.74], "content": [{"type": "text", "content": "图 1 营收"}]},
            {"type": "chart_footnote", "bbox": [0.08, 0.75, 0.91, 0.77],
             "content": [{"type": "text", "content": "从图 1 可以看出，西南区全年增速达到"}]},
        ]}
        elements = mineru_middle_to_elements(_middle([
            _t("上文。", [0.08, 0.1, 0.91, 0.12]), chart, _t("53.3%，增长最快。", [0.08, 0.78, 0.6, 0.8]),
        ]))
        self.assertEqual(_kinds(elements), ["text", "image", "text"])
        self.assertEqual(elements[1].text, "图 1 营收")
        self.assertEqual(elements[2].text, "从图 1 可以看出，西南区全年增速达到53.3%，增长最快。")

    def test_nested_list_items_on_separate_lines(self):
        nested = {"type": "list", "content": [
            _t("- 全年营收 4,860 万元"),
            {"type": "list", "content": [_t("- 华东区环比增长 17.1%"), _t("- 西南区增速 53.3%")]},
            _t("- 回款周期变长"),
        ]}
        elements = mineru_middle_to_elements(_middle([nested]))
        self.assertEqual(elements[0].text, "- 全年营收 4,860 万元\n  - 华东区环比增长 17.1%\n  - 西南区增速 53.3%\n- 回款周期变长")

    def test_wrapped_lines_are_joined_by_geometry(self):
        elements = mineru_middle_to_elements(_middle(
            [
                _t("甲方：星河科技有限公司", [0.08, 0.10, 0.34, 0.12]),           # 短行、没写满：不是折行
                _t("乙方：云帆数据服务有限公司", [0.08, 0.13, 0.39, 0.15]),
                _t("乙方提供运维服务，服务期限自 2025 年 3 月 1 日", [0.08, 0.16, 0.76, 0.18]),   # 写满整栏
                _t("至 2026 年 2 月 28 日。", [0.08, 0.19, 0.47, 0.21]),
                _t("乙方负有保密义务，", [0.08, 0.22, 0.60, 0.24]),              # 以逗号结尾
                _t("保密期限为三年。", [0.08, 0.25, 0.38, 0.27]),
                _t("星河科技合同专用章", [0.71, 0.84, 0.82, 0.88]),             # 印章：孤立的文字块
                _t("违约金为合同总额的 10%，从逾期之日起按日累计，直至付清为止且不", [0.08, 0.90, 0.76, 0.92]),
                _t("¹ 注释", [0.08, 0.95, 0.3, 0.96], "page_footnote"),
            ],
            [
                _t("超过合同总额。", [0.08, 0.05, 0.3, 0.07]),                   # 接上一页最后一行（跳过脚注）
                _t("第四条 保密条款", [0.08, 0.10, 0.27, 0.12], continues_prev=True),   # MinerU 误标，不能接到印章上
            ],
        ))
        texts = [e.text for e in elements]
        self.assertIn("甲方：星河科技有限公司", texts)
        self.assertIn("乙方：云帆数据服务有限公司", texts)
        self.assertIn("乙方提供运维服务，服务期限自 2025 年 3 月 1 日至 2026 年 2 月 28 日。", texts)
        self.assertIn("乙方负有保密义务，保密期限为三年。", texts)
        self.assertIn("违约金为合同总额的 10%，从逾期之日起按日累计，直至付清为止且不超过合同总额。", texts)
        self.assertIn("星河科技合同专用章", texts)
        self.assertIn("第四条 保密条款", texts)
        joined = next(e for e in elements if e.text.startswith("违约金"))
        self.assertEqual((joined.page, joined.extra["page_end"]), (1, 2))

    def test_office_paragraphs_without_bbox_are_not_joined(self):
        elements = mineru_middle_to_elements(_middle([_t("第一段没有句号"), _t("第二段")]))
        self.assertEqual([e.text for e in elements], ["第一段没有句号", "第二段"])

    def test_table_continued_on_next_page(self):
        def table(rows: str, **extra) -> dict:
            return {"type": "table", "content": [{"type": "table_body", "content": f"<table>{rows}</table>"}], **extra}

        head = "<tr><td>区域</td><td>营收</td></tr>"
        expected = "| 区域 | 营收 |\n| --- | --- |\n| 华东 | 480 |\n| 华南 | 390 |"
        for continued_rows in (head + "<tr><td>华南</td><td>390</td></tr>",    # 续页重复了表头
                               "<tr><td>华南</td><td>390</td></tr>"):          # 续页直接接数据行
            elements = mineru_middle_to_elements(_middle(
                [table(head + "<tr><td>华东</td><td>480</td></tr>")],
                [table(continued_rows, continues_prev=True)],
            ))
            self.assertEqual(len(elements), 1)
            self.assertEqual(elements[0].text, expected)
            self.assertEqual(elements[0].extra["page_end"], 2)


@registry.register("parser", "_test_broken")
def _broken_parser(**_kwargs):
    class Broken:
        def parse(self, path):
            raise RuntimeError("boom")
    return Broken()


class TestFileSource(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "docs"
        (self.root / "sub").mkdir(parents=True)
        (self.root / "a.md").write_text("# 标题\n\n内容一", encoding="utf-8")
        (self.root / "sub" / "b.html").write_text("<h1>网页</h1><p>内容二</p>", encoding="utf-8")
        (self.root / "c.bin").write_bytes(b"\x00\x01")
        (self.root / "empty.txt").write_text("   ", encoding="utf-8")
        self.cache = Path(self.tmp.name) / "cache"

    def tearDown(self):
        self.tmp.cleanup()

    def test_routes_and_metadata(self):
        docs = list(FileSource(str(self.root), cache_dir=str(self.cache)).load())
        self.assertEqual([d.doc_id for d in docs], ["a.md", "sub/b.html"])
        self.assertEqual(docs[0].text, "# 标题\n\n内容一")
        self.assertEqual((docs[1].metadata["parser"], docs[1].metadata["file_type"]), ("html", "html"))
        self.assertEqual(len(docs[0].metadata["file_hash"]), 64)

    def test_cache_skips_second_parse(self):
        list(FileSource(str(self.root), cache_dir=str(self.cache)).load())
        with patch("learn_rag.parsing.markdown.markdown_to_elements", side_effect=AssertionError("不应再解析")):
            docs = list(FileSource(str(self.root), cache_dir=str(self.cache)).load())
        self.assertEqual(docs[0].elements[0], Element("heading", "标题", level=1))

    def test_fallback_to_next_parser(self):
        source = FileSource(str(self.root / "a.md"), parsers={".md": ["_test_broken", "markdown"]}, cache_dir=None)
        with self.assertLogs("learn_rag.parsing.source", level="WARNING"):
            docs = list(source.load())
        self.assertEqual((docs[0].doc_id, docs[0].metadata["parser"]), ("a.md", "markdown"))

    def test_bad_ocr_config_does_not_break_pdf(self):
        try:
            import pymupdf
        except ImportError:
            self.skipTest("未安装 pymupdf")
        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 100), "A normal text page with enough characters to parse.", fontsize=11)
        pdf.save(self.root / "d.pdf")
        with patch.dict("os.environ", {"OCR_MODEL": ""}):
            source = FileSource(str(self.root / "d.pdf"), ocr={"type": "vlm_ocr"}, cache_dir=None)
            with self.assertLogs("learn_rag.parsing.source", level="WARNING"):
                docs = list(source.load())
        self.assertIn("normal text page", docs[0].text)


class TestTabular(unittest.TestCase):
    """CSV / TSV：编码兼容、表格输出、行数上限。"""

    def test_csv_utf8_bom(self):
        from learn_rag.parsing.tabular import CsvParser

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.csv"
            path.write_text("\ufeff区域,数量\n华东,120\n", encoding="utf-8")
            elements = CsvParser().parse(path)
        self.assertEqual([e.kind for e in elements], ["heading", "table"])
        self.assertEqual(elements[0].text, "orders")
        self.assertIn("| 区域 | 数量 |", elements[1].text)
        self.assertIn("| 华东 | 120 |", elements[1].text)

    def test_csv_gbk_fallback(self):
        from learn_rag.parsing.tabular import CsvParser

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "gbk.csv"
            path.write_text("区域,数量\n华北,80\n", encoding="gbk")
            elements = CsvParser().parse(path)
        self.assertIn("| 华北 | 80 |", elements[1].text)

    def test_tsv(self):
        from learn_rag.parsing.tabular import TsvParser

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "log.tsv"
            path.write_text("a\tb\n1\t2\n", encoding="utf-8")
            elements = TsvParser().parse(path)
        self.assertIn("| a | b |", elements[1].text)
        self.assertIn("| 1 | 2 |", elements[1].text)

    def test_max_rows_truncates(self):
        from learn_rag.parsing.tabular import CsvParser

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "big.csv"
            path.write_text("h\n" + "\n".join(str(i) for i in range(5)), encoding="utf-8")
            with self.assertLogs("learn_rag.parsing.tabular", level="WARNING"):
                elements = CsvParser(max_rows=2).parse(path)
        # 保留的 2 行 = 表头 + 1 行数据，Markdown 共 3 行
        self.assertEqual(len(elements[1].text.splitlines()), 3)


class TestLibreOffice(unittest.TestCase):
    """旧版 Office：转换命令、并发隔离配置、失败提示。"""

    def test_converts_and_delegates(self):
        from learn_rag.parsing.legacy import LibreOfficeParser

        seen = {}

        class FakeParser:
            def parse(self, path):
                seen["converted"] = path.name
                return [Element("text", "转换后的内容")]

        class FakeRegistry:
            def build(self, namespace, spec, **_):
                seen["spec"] = dict(spec)
                return FakeParser()

        def fake_run(cmd, **_):
            seen["cmd"] = cmd
            out = Path(cmd[cmd.index("--outdir") + 1])
            target = cmd[cmd.index("--convert-to") + 1]
            out.joinpath(f"legacy.{target}").write_bytes(b"")
            return SimpleNamespace(returncode=0)

        with patch("learn_rag.parsing.legacy.subprocess.run", side_effect=fake_run), \
             patch("learn_rag.parsing.legacy.registry", FakeRegistry()):
            elements = LibreOfficeParser().parse(Path("legacy.doc"))
        self.assertEqual([e.text for e in elements], ["转换后的内容"])
        self.assertEqual(seen["converted"], "legacy.docx")
        self.assertEqual(seen["spec"], {"type": "docx"})
        self.assertIn("--headless", seen["cmd"])
        # 并发转换必须使用独立的用户配置目录，否则会抢 LibreOffice 的配置锁
        self.assertTrue(any(a.startswith("-env:UserInstallation=") for a in seen["cmd"]))

    def test_missing_command_message(self):
        from learn_rag.parsing.legacy import LibreOfficeParser

        with patch("learn_rag.parsing.legacy.subprocess.run", side_effect=FileNotFoundError):
            with self.assertRaisesRegex(RuntimeError, "LibreOffice"):
                LibreOfficeParser().parse(Path("a.doc"))

    def test_unsupported_suffix(self):
        from learn_rag.parsing.legacy import LibreOfficeParser

        with self.assertRaises(ValueError):
            LibreOfficeParser().parse(Path("a.docx"))


class TestFileSourceHardening(unittest.TestCase):
    """批处理稳定性：坏文件只能被隔离，不能中断整批导入。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_skips_temp_and_hidden_files(self):
        (self.root / "a.md").write_text("# 好\n\n正文", encoding="utf-8")
        (self.root / "~$a.docx").write_bytes(b"lock")          # Office 打开文档时的锁文件
        (self.root / ".hidden.md").write_text("# 隐藏", encoding="utf-8")
        (self.root / "b.tmp").write_text("x", encoding="utf-8")
        docs = list(FileSource(str(self.root), cache_dir=None).load())
        self.assertEqual([d.doc_id for d in docs], ["a.md"])

    def test_max_file_size(self):
        (self.root / "small.md").write_text("# 小\n\n正文", encoding="utf-8")
        (self.root / "big.md").write_text("# 大\n\n" + "字" * 100, encoding="utf-8")
        docs = list(FileSource(str(self.root), cache_dir=None, max_file_size=50).load())
        self.assertEqual([d.doc_id for d in docs], ["small.md"])

    def test_unreadable_file_is_isolated(self):
        (self.root / "good.md").write_text("# 好\n\n正文", encoding="utf-8")
        (self.root / "bad.md").write_text("# 坏", encoding="utf-8")
        original = Path.read_bytes

        def fake_read_bytes(path):
            if path.name == "bad.md":
                raise OSError("permission denied")
            return original(path)

        with patch.object(Path, "read_bytes", fake_read_bytes):
            with self.assertLogs("learn_rag.parsing.source", level="ERROR"):
                docs = list(FileSource(str(self.root), cache_dir=None).load())
        self.assertEqual([d.doc_id for d in docs], ["good.md"])

    def test_corrupt_cache_is_rebuilt(self):
        (self.root / "a.md").write_text("# 标题\n\n正文", encoding="utf-8")
        cache = self.root / "cache"
        list(FileSource(str(self.root), cache_dir=str(cache)).load())
        cache_file = next(cache.glob("*.json"))
        cache_file.write_text("{broken", encoding="utf-8")

        with self.assertLogs("learn_rag.parsing.source", level="WARNING"):
            docs = list(FileSource(str(self.root), cache_dir=str(cache)).load())
        self.assertEqual(docs[0].elements[0], Element("heading", "标题", level=1))
        self.assertEqual(json.loads(cache_file.read_text(encoding="utf-8"))[0]["kind"], "heading")

    def test_legacy_office_missing_soffice_does_not_break_batch(self):
        (self.root / "a.md").write_text("# 好\n\n正文", encoding="utf-8")
        (self.root / "old.doc").write_bytes(b"ole2")
        with patch("learn_rag.parsing.legacy.subprocess.run", side_effect=FileNotFoundError):
            with self.assertLogs("learn_rag.parsing.source", level="ERROR"):
                docs = list(FileSource(str(self.root), cache_dir=None).load())
        self.assertEqual([d.doc_id for d in docs], ["a.md"])


class TestEnterpriseRoutes(unittest.TestCase):
    """企业级路由：PDF 首选 MinerU、Word 首选 Docling，未安装时自动降级且不中断。"""

    def test_config_routes(self):
        from learn_rag.core.config import load_config

        cfg = load_config("configs/parsing.yaml")
        self.assertEqual(cfg["parsing"]["parsers"][".pdf"], ["mineru", "pdf"])
        self.assertEqual(cfg["parsing"]["parsers"][".docx"], ["docling", "docx"])

    def test_pdf_falls_back_to_pymupdf(self):
        try:
            import pymupdf  # noqa: F401
        except ImportError:
            self.skipTest("未安装 pymupdf")
        from learn_rag.parsing.external import MinerUParser

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf = pymupdf.open()
            pdf.new_page().insert_text((72, 100), "A normal text page with enough characters.", fontsize=11)
            pdf.save(root / "a.pdf")
            with patch.object(MinerUParser, "parse", side_effect=RuntimeError("找不到 MinerU 命令")):
                with self.assertLogs("learn_rag.parsing.source", level="WARNING"):
                    docs = list(FileSource(str(root), parsers={".pdf": ["mineru", "pdf"]}, cache_dir=None).load())
        self.assertEqual(docs[0].metadata["parser"], "pdf")
        self.assertIn("normal text page", docs[0].text)

    def test_docx_falls_back_to_python_docx(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("未安装 python-docx")
        from learn_rag.parsing.external import DoclingParser

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            d = docx.Document()
            d.add_heading("制度", level=1)
            d.add_paragraph("一线城市每晚上限 600 元。")
            d.save(root / "a.docx")
            with patch.object(DoclingParser, "parse", side_effect=RuntimeError("未安装 Docling")):
                with self.assertLogs("learn_rag.parsing.source", level="WARNING"):
                    docs = list(FileSource(str(root), parsers={".docx": ["docling", "docx"]}, cache_dir=None).load())
        self.assertEqual(docs[0].metadata["parser"], "docx")
        self.assertIn("600 元", docs[0].text)


class TestAssetStorage(unittest.TestCase):
    """图片资产：内容寻址去重、MinerU/Docling 提取、路径进入 chunk 元数据。"""

    def test_save_asset_dedupes(self):
        from learn_rag.parsing.assets import save_asset

        with tempfile.TemporaryDirectory() as tmp:
            first = save_asset(b"png-bytes", tmp, ext="png", mime="image/png")
            second = save_asset(b"png-bytes", tmp, ext="png", mime="image/png")
            self.assertEqual(first, second)
            self.assertEqual((Path(tmp) / first["asset"]).read_bytes(), b"png-bytes")
            self.assertEqual(len(list(Path(tmp).rglob("*.png"))), 1)

    def test_mineru_persists_images(self):
        import hashlib

        def fake_run(cmd, **_):
            TestExternalAdapters._write_mineru_zip(Path(cmd[cmd.index("-o") + 1]))
            return SimpleNamespace(returncode=0)

        with tempfile.TemporaryDirectory() as tmp, \
             patch("learn_rag.parsing.external.subprocess.run", side_effect=fake_run):
            elements = MinerUParser(assets_dir=tmp).parse(Path("a.pdf"))
            for element in elements:
                if element.kind == "table":
                    expected = TestExternalAdapters.IMAGES["images/table.jpg"]
                elif element.kind == "image":
                    name = {"image": "pic", "chart": "chart"}[element.extra["mineru_type"]]
                    expected = TestExternalAdapters.IMAGES[f"images/{name}.jpg"]
                else:
                    continue
                self.assertEqual(element.extra["sha256"], hashlib.sha256(expected).hexdigest())
                self.assertEqual((Path(tmp) / element.extra["asset"]).read_bytes(), expected)

    def test_mineru_rejects_unsafe_image_path(self):
        unsafe = {
            "schema": "docvortex.middle",
            "pages": [{
                "page_idx": 0,
                "blocks": [{
                    "type": "image",
                    "content": [{"type": "image_body", "image_path": "../evil.jpg", "content": ""}],
                }],
            }],
        }

        def fake_run(cmd, **_):
            TestExternalAdapters._write_mineru_zip(
                Path(cmd[cmd.index("-o") + 1]), middle=unsafe, images={"evil.jpg": b"bad"}
            )
            return SimpleNamespace(returncode=0)

        with tempfile.TemporaryDirectory() as tmp, \
             patch("learn_rag.parsing.external.subprocess.run", side_effect=fake_run), \
             self.assertLogs("learn_rag.parsing.external", level="WARNING"):
            elements = MinerUParser(assets_dir=tmp).parse(Path("a.pdf"))
        self.assertEqual((elements[0].text, elements[0].extra), ("[图片]", {"mineru_type": "image"}))

    def test_docling_persists_pictures(self):
        class FakePil:
            def save(self, fh, format=None, **_):
                fh.write(b"png-bytes")

        picture = SimpleNamespace(
            label=SimpleNamespace(value="picture"),
            text="",
            prov=[SimpleNamespace(page_no=1, bbox=SimpleNamespace(l=0, t=1, r=2, b=3))],
            caption_text=lambda doc: "图1 流程图",
            image=SimpleNamespace(pil_image=FakePil(), mimetype="image/png"),
        )
        doc = SimpleNamespace(iterate_items=lambda: [(picture, 0)])
        with tempfile.TemporaryDirectory() as tmp:
            elements = docling_to_elements(doc, assets_dir=tmp)
            self.assertEqual(elements[0].text, "图1 流程图")
            self.assertEqual((Path(tmp) / elements[0].extra["asset"]).read_bytes(), b"png-bytes")

    def test_image_asset_reaches_chunk_metadata(self):
        doc = Document.from_elements("manual.pdf", [
            Element("heading", "架构", level=1),
            Element("image", "图1 总体架构", extra={"asset": "ab/abc.png", "sha256": "abc", "mime": "image/png"}),
        ], {"title": "手册"})
        chunks = StructureChunker(chunk_size=200).split(doc)
        self.assertEqual(chunks[0].metadata["assets"],
                         [{"asset": "ab/abc.png", "kind": "image", "caption": "图1 总体架构", "mime": "image/png"}])


class TestEndToEnd(unittest.TestCase):
    def test_parse_chunk_answer(self):
        from learn_rag.core.config import load_config
        from learn_rag.pipeline.rag import RagPipeline

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "policy.md").write_text(
                "# 报销制度\n\n## 住宿标准\n\n一线城市每晚上限 600 元。\n\n## 餐饮补贴\n\n每人每天 120 元。",
                encoding="utf-8",
            )
            (root / "faq.html").write_text("<h1>常见问题</h1><p>发票需在 30 天内提交。</p>", encoding="utf-8")
            cfg = load_config("configs/default.yaml", "configs/parsing.yaml")
            cfg["parsing"]["cache_dir"] = str(root / "cache")
            pipe = RagPipeline.from_config(cfg)
            stats = pipe.index(FileSource(str(root), **cfg["parsing"]).load())
            result = pipe.answer("餐饮补贴每天多少钱？")
        self.assertEqual(stats["documents"], 2)
        self.assertIn("120", result.answer.text)
        self.assertEqual(result.contexts[0].chunk.metadata["section"], "报销制度 > 餐饮补贴")


if __name__ == "__main__":
    unittest.main()
