"""
server.recommend —— "一键推荐"：扫描上传的文件，按规则给出解析器、切分参数和 embedding 模型，每一项附上理由。

用规则而不是实测择优：规则秒出结果、每条都说得清为什么，适合学习；实测择优（自动出题、
多组参数各建一次库比召回率）需要大模型、耗时几十倍，留到以后做成"验证推荐"。

扫描只做轻量的统计，不完整解析文档：PDF 抽样几页看有没有文字层、表格和图片，Office 读结构计数，
文本文件读正文。任何一个文件读不了只是少一份统计，不影响推荐。
"""

from __future__ import annotations

import logging
import re
from collections import Counter
from pathlib import Path
from typing import Any

from ..core.text import count_tokens
from . import environment

logger = logging.getLogger(__name__)

_TYPE_NAMES = {
    "pdf": "PDF", "docx": "Word", "doc": "Word", "pptx": "PPT", "ppt": "PPT", "xlsx": "Excel", "xls": "Excel",
    "csv": "CSV", "tsv": "TSV", "md": "Markdown", "markdown": "Markdown", "txt": "文本", "html": "网页",
    "htm": "网页", "rtf": "RTF", "png": "图片", "jpg": "图片", "jpeg": "图片", "webp": "图片",
}
_TABULAR = {"xlsx", "xls", "csv", "tsv"}
_IMAGES = {"png", "jpg", "jpeg", "webp"}
_PDF_SAMPLE_PAGES = 30      # PDF 最多抽样的页数，统计结果按比例放大
_PDF_TABLE_PAGES = 5        # 表格检测比较慢，只在其中几页上做
_CJK = re.compile(r"[一-鿿]")
_LATIN = re.compile(r"[A-Za-z]")


def profile(files: list[Path]) -> dict[str, Any]:
    """统计一批文件：格式分布、页数、扫描页、表格、图片、标题、代码块、篇幅、中英文比例、长段落比例。"""
    total: dict[str, Any] = {
        "files": len(files), "types": Counter(), "pages": 0, "scanned_pages": 0, "tables": 0, "images": 0,
        "headings": 0, "code_blocks": 0, "tokens": 0, "paragraphs": [], "unreadable": [],
    }
    cjk = latin = 0
    for file in files:
        kind = file.suffix.lower().lstrip(".")
        total["types"][kind] += 1
        try:
            stats = _scan(file, kind)
        except Exception as exc:   # 扫描只是为了推荐，一个文件读不了不影响其他文件
            logger.warning("扫描 %s 失败：%s", file.name, exc)
            total["unreadable"].append(file.name)
            continue
        text = stats.pop("text", "")
        scale = stats.pop("scale", 1.0)
        total["tokens"] += int(count_tokens(text) * scale)
        cjk += len(_CJK.findall(text))
        latin += len(_LATIN.findall(text))
        total["paragraphs"] += stats.pop("paragraphs", [])
        for key, value in stats.items():
            total[key] += value
    total["types"] = dict(total["types"])
    total["cjk_ratio"] = round(cjk / (cjk + latin / 5), 2) if cjk + latin else 0.0   # 英文按约 5 个字母一个词折算
    return total


def _scan(file: Path, kind: str) -> dict[str, Any]:
    if kind == "pdf":
        return _scan_pdf(file)
    if kind == "docx":
        return _scan_docx(file)
    if kind == "pptx":
        return _scan_pptx(file)
    if kind == "xlsx":
        return _scan_xlsx(file)
    if kind in ("csv", "tsv"):
        lines = file.read_text(encoding="utf-8", errors="replace").splitlines()
        return {"tables": 1, "text": "\n".join(lines[:200]), "scale": max(1.0, len(lines) / 200)}
    if kind in _IMAGES:
        return {"images": 1, "scanned_pages": 1, "pages": 1}
    if kind in ("md", "markdown", "txt"):
        return _scan_markdown(file.read_text(encoding="utf-8", errors="replace"))
    if kind in ("html", "htm"):
        return _scan_html(file.read_text(encoding="utf-8", errors="replace"))
    return {}   # .doc / .xls / .ppt / .rtf 要先转换格式，这里只计数


def _scan_pdf(file: Path) -> dict[str, Any]:
    import pymupdf  # 延迟导入

    with pymupdf.open(file) as doc:
        pages = doc.page_count
        step = max(1, pages // _PDF_SAMPLE_PAGES)
        sampled = list(range(0, pages, step))[:_PDF_SAMPLE_PAGES]
        scanned = images = tables = 0
        texts = []
        for n, index in enumerate(sampled):
            page = doc[index]
            text = page.get_text()
            texts.append(text)
            if len(text.strip()) < 20:     # 几乎没有文字层：扫描页或纯图片页
                scanned += 1
            images += len(page.get_images())
            if n < _PDF_TABLE_PAGES and hasattr(page, "find_tables"):
                try:
                    tables += len(page.find_tables().tables)
                except Exception:
                    pass
        scale = pages / max(1, len(sampled))
        table_scale = pages / max(1, min(len(sampled), _PDF_TABLE_PAGES))
    # PDF 的换行是排版换行，分不出段落，不参与长段落统计
    return {"pages": pages, "scanned_pages": round(scanned * scale), "images": round(images * scale),
            "tables": round(tables * table_scale), "text": "\n".join(texts), "scale": scale}


def _scan_docx(file: Path) -> dict[str, Any]:
    import docx  # python-docx，延迟导入

    document = docx.Document(str(file))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    headings = sum(1 for p in paragraphs if (p.style.name or "").lower().startswith(("heading", "title", "标题")))
    body = [p.text for p in paragraphs]
    return {"headings": headings, "tables": len(document.tables), "images": len(document.inline_shapes),
            "text": "\n\n".join(body), "paragraphs": [count_tokens(t) for t in body]}


def _scan_pptx(file: Path) -> dict[str, Any]:
    from pptx import Presentation  # python-pptx，延迟导入

    deck = Presentation(str(file))
    texts, tables, images = [], 0, 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                tables += 1
            if shape.shape_type == 13:   # MSO_SHAPE_TYPE.PICTURE
                images += 1
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                texts.append(shape.text_frame.text)
    slides = len(deck.slides)
    return {"pages": slides, "headings": slides, "tables": tables, "images": images, "text": "\n".join(texts)}


def _scan_xlsx(file: Path) -> dict[str, Any]:
    from openpyxl import load_workbook  # 延迟导入

    workbook = load_workbook(str(file), read_only=True, data_only=True)
    try:
        rows, texts = 0, []
        for sheet in workbook.worksheets:
            for n, row in enumerate(sheet.iter_rows(values_only=True)):
                rows += 1
                if n < 50:
                    texts.append(" ".join(str(v) for v in row if v is not None))
        sampled = max(1, len(texts))
        return {"tables": len(workbook.worksheets), "text": "\n".join(texts), "scale": max(1.0, rows / sampled)}
    finally:
        workbook.close()


def _scan_markdown(text: str) -> dict[str, Any]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    prose = [p for p in paragraphs if not p.startswith(("#", "|", "```", "!["))]
    return {
        "headings": len(re.findall(r"^#{1,6}\s", text, re.M)),
        "tables": len(re.findall(r"^\|?\s*:?-{3,}", text, re.M)),
        "images": text.count("!["),
        "code_blocks": text.count("```") // 2,
        "text": text,
        "paragraphs": [count_tokens(p) for p in prose],
    }


def _scan_html(html: str) -> dict[str, Any]:
    text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    return {
        "headings": len(re.findall(r"<h[1-6][\s>]", html, re.I)),
        "tables": len(re.findall(r"<table[\s>]", html, re.I)),
        "images": len(re.findall(r"<img[\s>]", html, re.I)),
        "code_blocks": len(re.findall(r"<pre[\s>]", html, re.I)),
        "text": text,
    }


def recommend(stats: dict[str, Any]) -> dict[str, Any]:
    """
    按统计结果推荐参数，返回 {summary, settings, reasons, warnings}。
    reasons 的键是表单字段（chunker / chunk_size / parent_size / overlap_sentences / encoder / index /
    pdf_parser / docx_parser / ocr），值是一句话理由，前端显示在对应字段旁边。
    """
    settings = environment.default_settings()
    reasons: dict[str, str] = {}
    warnings: list[str] = []
    files = max(1, stats["files"])
    types = stats["types"]
    avg_tokens = stats["tokens"] / files
    tabular_share = sum(types.get(t, 0) for t in _TABULAR) / files
    table_density = stats["tables"] / max(1.0, stats["tokens"] / 2000)   # 每 2000 token 有几张表

    # 切分器：结构切分对有结构和没结构的文本都适用，没有理由换
    reasons["chunker"] = "按标题和段落结构切分：表格、图片和说明文字不会被切散；纯文本也会按空行分段走同一套规则"

    # 子块大小：表格和代码按行切，块小了会被切得很碎，所以先看它们，再看篇幅
    if tabular_share >= 0.3 or table_density >= 1:
        chunk_size = 400
        reasons["chunk_size"] = (f"表格较多（{stats['tables']} 张）：表格按行切，每片都要重复表题和表头，"
                                 "块大一些重复更少、每片装的行更多")
    elif stats["code_blocks"] >= 5:
        chunk_size = 400
        reasons["chunk_size"] = f"代码块较多（{stats['code_blocks']} 个）：代码按行切，块大一些函数不容易被拦腰切断"
    elif avg_tokens and avg_tokens < 600:
        chunk_size = 200
        reasons["chunk_size"] = f"文档普遍较短（平均约 {avg_tokens:.0f} token），小块让每块只讲一件事，召回更准"
    else:
        chunk_size = 300
        reasons["chunk_size"] = "通用默认：约 300 token 一块（中文约 300 字），语义集中又不至于丢上下文"

    # 父块大小
    parent_size = chunk_size * 4
    if avg_tokens and avg_tokens < chunk_size * 2:
        reasons["parent_size"] = (f"约为子块的 4 倍。文档平均只有 {avg_tokens:.0f} token，多数文档只切出一个子块，"
                                  "这时不会生成父块，设多大都不影响")
    else:
        reasons["parent_size"] = "约为子块的 4 倍，大致是一个完整小节：命中子块后交给大模型的上下文更完整"

    # 句子重叠
    paragraphs = stats["paragraphs"]
    long_share = sum(1 for p in paragraphs if p > chunk_size) / len(paragraphs) if paragraphs else 0.0
    if long_share >= 0.2:
        overlap = 1
        reasons["overlap_sentences"] = (f"{long_share:.0%} 的段落比子块还长，切口会落在段落中间；"
                                        "相邻子块重叠 1 句，避免“它”“该方案”这类指代在边界处丢失")
    else:
        overlap = 0
        reasons["overlap_sentences"] = "段落大多比子块短，切口都在段落或句子边界上，上下文由父块补全，不需要重叠"

    settings["chunker"] = {"type": "structure", "chunk_size": chunk_size, "parent_size": parent_size,
                           "overlap_sentences": overlap}

    # embedding 模型：只在能用的里面选
    if settings["encoder"]["type"] == "hashing":
        reasons["encoder"] = "没有配置 embedding 服务，先用 Hashing 跑通流程；它不理解语义，配置好服务后建议换模型重建"
        warnings.append("当前只能使用 Hashing 向量，检索效果只用于演示。在 .env 配置 EMBEDDING_* 后可选真实模型")
    else:
        reasons["encoder"] = f"已配置 embedding 服务（{settings['encoder']['model']}），语义检索效果远好于 Hashing"
    cjk = stats["cjk_ratio"]
    if stats["tokens"]:
        language = "中文为主" if cjk >= 0.6 else "英文为主" if cjk <= 0.2 else "中英混合"
        reasons["encoder"] += f"；文档{language}，选模型时注意它对{'中文' if cjk >= 0.2 else '英文'}的支持"

    # 向量库
    if settings["index"]["type"] == "chroma":
        reasons["index"] = "Chroma 写入即落盘，支持增量更新和元数据过滤"
    else:
        reasons["index"] = "未安装 chromadb，使用内存索引（flat），建完保存到磁盘；数据量大时建议安装 Chroma"

    # 解析器
    scanned = stats["scanned_pages"]
    if types.get("pdf"):
        if environment.has_mineru() and (scanned or stats["tables"]):
            settings["parsers"][".pdf"] = ["mineru", "pdf"]
            found = f"有 {scanned} 页扫描页" if scanned else "PDF 里有表格"
            reasons["pdf_parser"] = found + "，MinerU 的版面、表格和 OCR 更准；它失败时自动退回 PyMuPDF"
        elif scanned:
            settings["parsers"][".pdf"] = ["pdf"]
            reasons["pdf_parser"] = f"有 {scanned} 页扫描页，但没有安装 MinerU，只能交给 OCR 服务识别"
        else:
            settings["parsers"][".pdf"] = ["pdf"]
            reasons["pdf_parser"] = "PDF 都有文字层，PyMuPDF 直接读取又快又准；版面复杂时可换 MinerU"
    if types.get("docx"):
        reasons["docx_parser"] = ("Docling 对 Word 的标题、表格和图片还原更完整，失败时退回 python-docx"
                                  if settings["parsers"][".docx"][0] == "docling"
                                  else "python-docx 直接读取 Word 的标题、段落和表格，速度快")
    images = sum(types.get(t, 0) for t in _IMAGES)
    if scanned or images:
        settings["ocr"] = environment.has_ocr()
        if settings["ocr"]:
            reasons["ocr"] = "有扫描页或图片文件，需要 OCR 才能读出文字"
        else:
            reasons["ocr"] = "有扫描页或图片文件，但没有配置 OCR 服务（OCR_BASE_URL），这些页面会被跳过"
            warnings.append(f"{scanned} 页扫描页没有可用的 OCR，内容会缺失。配置 OCR_* 或安装 MinerU 后重建")

    if stats["unreadable"]:
        warnings.append(f"{len(stats['unreadable'])} 个文件无法预览，未计入统计：{'、'.join(stats['unreadable'][:3])}")
    return {"summary": _summary(stats), "settings": settings, "reasons": reasons, "warnings": warnings}


def _summary(stats: dict[str, Any]) -> list[str]:
    names: Counter = Counter()
    for kind, n in stats["types"].items():
        names[_TYPE_NAMES.get(kind, kind.upper() or "其他")] += n
    parts = []
    for name, n in names.most_common():
        extra = f"（扫描页 {stats['scanned_pages']}）" if name == "PDF" and stats["scanned_pages"] else ""
        parts.append(f"{name} {n}{extra}")
    lines = [f"{stats['files']} 个文件：" + "、".join(parts)]
    if stats["tokens"]:
        lines.append(f"约 {_short(stats['tokens'])} token，平均每篇 {_short(stats['tokens'] / max(1, stats['files']))}")
    lines.append(f"表格 {stats['tables']} · 图片 {stats['images']} · 标题 {stats['headings']}")
    if stats["tokens"]:
        cjk = stats["cjk_ratio"]
        lines.append("中文为主" if cjk >= 0.6 else "英文为主" if cjk <= 0.2 else "中英混合")
    return lines


def _short(n: float) -> str:
    return f"{n / 10000:.1f} 万" if n >= 10000 else f"{n:.0f}"
