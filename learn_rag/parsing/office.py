"""
parsing.office —— Word / PowerPoint / Excel 解析。

Office 文件本身就保存了结构（标题样式、表格、幻灯片），直接读取即可，不需要 OCR，
既快又准。三个解析器各自只负责把一种格式的结构翻译成 Element。

依赖（pip install -e ".[parsing]"）在 parse 时才导入，没装也不影响其他格式。
"""

from __future__ import annotations

import re
from pathlib import Path

from ..core.interfaces import DocumentParser
from ..core.registry import registry
from ..core.types import Element
from .markdown import table_to_markdown

# Word 的标题样式名：英文版 "Heading 2"，中文版 "标题 2"，文档标题 "Title"
_HEADING_STYLE = re.compile(r"(?:heading|标题)\s*(\d)", re.I)


@registry.register("parser", "docx")
class DocxParser(DocumentParser):
    """
    Word（.docx）：按正文顺序读取段落和表格，标题层级取自段落样式。

    python-docx 的 doc.paragraphs 与 doc.tables 是分开的两个列表，会丢失二者的先后顺序，
    所以这里直接遍历正文的 XML 子节点，保证表格出现在它原本的位置。
    """

    def parse(self, path: Path) -> list[Element]:
        from docx import Document as open_docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = open_docx(str(path))
        elements: list[Element] = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                if text:
                    level = _heading_level(paragraph.style.name if paragraph.style is not None else "")
                    elements.append(Element("heading", text, level=level) if level else Element("text", text))
            elif tag == "tbl":
                table = Table(child, doc)
                markdown = table_to_markdown([[cell.text for cell in row.cells] for row in table.rows])
                if markdown:
                    elements.append(Element("table", markdown))
        return elements


def _heading_level(style_name: str) -> int:
    if style_name.lower() == "title":
        return 1
    match = _HEADING_STYLE.search(style_name)
    return int(match.group(1)) if match else 0


@registry.register("parser", "pptx")
class PptxParser(DocumentParser):
    """
    PowerPoint（.pptx）：每页幻灯片是一个一级标题，页内形状按从上到下、从左到右的阅读顺序输出。

    演讲者备注往往比幻灯片本身信息更完整，一并收录。page 记录幻灯片页码，便于回答时溯源。
    """

    def parse(self, path: Path) -> list[Element]:
        from pptx import Presentation

        elements: list[Element] = []
        for number, slide in enumerate(Presentation(str(path)).slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text_frame.text.strip() if title_shape is not None else ""
            elements.append(Element("heading", title or f"第 {number} 页", level=1, page=number))

            for shape in sorted(_flatten_shapes(slide.shapes), key=lambda s: (s.top or 0, s.left or 0)):
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if shape.has_text_frame and shape.text_frame.text.strip():
                    elements.append(Element("text", shape.text_frame.text.strip(), page=number))
                elif getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    markdown = table_to_markdown(rows)
                    if markdown:
                        elements.append(Element("table", markdown, page=number))

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    elements.append(Element("text", f"备注：{notes}", page=number, extra={"notes": True}))
        return elements


def _flatten_shapes(shapes) -> list:
    """组合形状里还嵌着形状，展开成一层，调用方就不用处理嵌套。"""
    out = []
    for shape in shapes:
        if hasattr(shape, "shapes"):
            out.extend(_flatten_shapes(shape.shapes))
        else:
            out.append(shape)
    return out


@registry.register("parser", "xlsx")
class XlsxParser(DocumentParser):
    """
    Excel（.xlsx）：每个工作表是一个一级标题加一张表格，第一行视为表头。

    读取的是公式的计算结果而不是公式本身（data_only=True）。
    全空的行和列直接去掉 —— Excel 的"已用区域"经常比真实数据大得多。
    超长的表格交给结构切分器按行切分，并在每块重复表头。
    """

    def parse(self, path: Path) -> list[Element]:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        elements: list[Element] = []
        try:
            for sheet in workbook.worksheets:
                rows = _drop_empty_columns([
                    ["" if v is None else str(v) for v in row]
                    for row in sheet.iter_rows(values_only=True)
                    if any(v is not None and str(v).strip() for v in row)
                ])
                markdown = table_to_markdown(rows)
                if markdown:
                    elements.append(Element("heading", sheet.title, level=1))
                    elements.append(Element("table", markdown, extra={"sheet": sheet.title}))
        finally:
            workbook.close()
        return elements


def _drop_empty_columns(rows: list[list[str]]) -> list[list[str]]:
    if not rows:
        return rows
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [j for j in range(width) if any(r[j].strip() for r in rows)]
    return [[r[j] for j in keep] for r in rows]
